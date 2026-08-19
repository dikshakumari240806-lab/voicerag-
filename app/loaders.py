from pathlib import Path
from pypdf import PdfReader
from docx import Document as DocxDocument
from .models import Document

def load_file(path):
    ext=path.suffix.lower()
    if ext in {".txt",".md"}:
        text=path.read_text(encoding="utf-8",errors="ignore")
    elif ext==".pdf":
        text="\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    elif ext==".docx":
        text="\n".join(p.text for p in DocxDocument(str(path)).paragraphs)
    else: raise ValueError(f"Unsupported: {path}")
    return Document(str(path),text.strip())

def load_directory(directory="data"):
    out=[]
    for p in sorted(Path(directory).rglob("*")):
        if p.is_file() and p.suffix.lower() in {".txt",".md",".pdf",".docx"}:
            d=load_file(p)
            if d.text: out.append(d)
    return out
